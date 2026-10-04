import os

from flask import Flask, render_template, request, redirect, url_for, session, flash, Response
from dotenv import load_dotenv
import sqlite3
import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from werkzeug.security import check_password_hash, generate_password_hash


app = Flask(__name__)

load_dotenv()

DATABASE_PATH = "database/attendance.db"

# Secret key for Flask sessions
app.secret_key = os.getenv("SECRET_KEY")


@app.route("/", methods=["GET", "POST"])
def login():

    if request.method == "POST":

        username = request.form["username"].strip()
        password = request.form["password"]

        connection = sqlite3.connect(DATABASE_PATH)
        connection.row_factory = sqlite3.Row

        user = connection.execute(
            "SELECT * FROM users WHERE username = ?",
            (username,)
        ).fetchone()

        connection.close()

        if user and check_password_hash(
            user["password_hash"],
            password
        ):
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user["role"]

            if user["role"] == "HOD":
                return redirect(url_for("hod_dashboard"))

            if user["role"] == "FACULTY":
                return redirect(url_for("faculty_dashboard"))

            if user["role"] == "STUDENT":
                return redirect(url_for("student_dashboard"))

            flash("Invalid user role.", "error")
            return redirect(url_for("login"))

        flash("Invalid username or password.", "error")

    return render_template("login.html")


@app.route("/hod/dashboard")
def hod_dashboard():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    today = __import__("datetime").date.today().isoformat()

    total_students = connection.execute("""
        SELECT COUNT(*) AS total
        FROM students
    """).fetchone()["total"]

    present_today = connection.execute("""
        SELECT COUNT(*) AS total
        FROM attendance
        WHERE date = ?
        AND status = 'PRESENT'
    """, (today,)).fetchone()["total"]

    absent_today = connection.execute("""
        SELECT COUNT(*) AS total
        FROM attendance
        WHERE date = ?
        AND status = 'ABSENT'
    """, (today,)).fetchone()["total"]

    total_attendance = connection.execute("""
        SELECT COUNT(*) AS total
        FROM attendance
    """).fetchone()["total"]

    if total_attendance > 0:
        attendance_percentage = round(
            (connection.execute("""
                SELECT COUNT(*) AS total
                FROM attendance
                WHERE status = 'PRESENT'
            """).fetchone()["total"] / total_attendance) * 100,
            2
        )
    else:
        attendance_percentage = 0

    connection.close()

    return render_template(
        "hod/dashboard.html",
        today=today,
        total_students=total_students,
        present_today=present_today,
        absent_today=absent_today,
        attendance_percentage=attendance_percentage
    )

@app.route("/hod/students")
def hod_students():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    search = request.args.get("search", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    if search:

        search_pattern = f"%{search}%"

        students = connection.execute("""
            SELECT
                id,
                student_id,
                name,
                email,
                branch,
                year,
                section
            FROM students
            WHERE
                student_id LIKE ?
                OR name LIKE ?
                OR branch LIKE ?
                OR section LIKE ?
            ORDER BY id DESC
        """, (
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern
        )).fetchall()

    else:

        students = connection.execute("""
            SELECT
                id,
                student_id,
                name,
                email,
                branch,
                year,
                section
            FROM students
            ORDER BY id DESC
        """).fetchall()

    connection.close()

    return render_template(
        "hod/students.html",
        students=students,
        search=search
    )

@app.route("/hod/students/add", methods=["GET", "POST"])
def add_student():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    if request.method == "POST":

        student_id = request.form["student_id"].strip()
        name = request.form["name"].strip()
        email = request.form["email"].strip()
        branch = request.form["branch"].strip()
        year = request.form["year"]
        section = request.form["section"]

        # Student input validation

        if not student_id:
            flash("Student ID is required.", "error")
            return redirect(url_for("add_student"))

        if not name:
            flash("Student name is required.", "error")
            return redirect(url_for("add_student"))

        if not email:
            flash("Email is required.", "error")
            return redirect(url_for("add_student"))

        if "@" not in email or "." not in email:
            flash("Please enter a valid email address.", "error")
            return redirect(url_for("add_student"))

        if not branch:
            flash("Branch is required.", "error")
            return redirect(url_for("add_student"))

        if year not in ["1", "2", "3", "4"]:
            flash("Invalid year selected.", "error")
            return redirect(url_for("add_student"))

        if section not in ["A", "B", "C", "D"]:
            flash("Invalid section selected.", "error")
            return redirect(url_for("add_student"))

        connection = sqlite3.connect(DATABASE_PATH)

        existing_student = connection.execute("""
            SELECT id
            FROM students
            WHERE student_id = ?
        """, (student_id,)).fetchone()

        if existing_student:
            connection.close()

            flash(
                "Student ID already exists.",
                "error"
            )

            return redirect(url_for("add_student"))

        connection.execute("""
            INSERT INTO students (
                student_id,
                name,
                email,
                branch,
                year,
                section
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            student_id,
            name,
            email,
            branch,
            year,
            section
        ))

        connection.commit()
        connection.close()

        flash(
            "Student added successfully.",
            "success"
        )

        return redirect(url_for("hod_students"))

    return render_template("hod/add_student.html")

@app.route("/hod/students/edit/<int:student_id>", methods=["GET", "POST"])
def edit_student(student_id):

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    student = connection.execute("""
        SELECT *
        FROM students
        WHERE id = ?
    """, (student_id,)).fetchone()

    if not student:
        connection.close()
        return "Student not found", 404

    if request.method == "POST":

        new_student_id = request.form["student_id"].strip()
        name = request.form["name"].strip()
        email = request.form["email"].strip()
        branch = request.form["branch"].strip()
        year = request.form["year"]
        section = request.form["section"]

        # Student input validation

        if not new_student_id:
            flash("Student ID is required.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if not name:
            flash("Student name is required.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if not email:
            flash("Email is required.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if "@" not in email or "." not in email:
            flash("Please enter a valid email address.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if not branch:
            flash("Branch is required.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if year not in ["1", "2", "3", "4"]:
            flash("Invalid year selected.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        if section not in ["A", "B", "C", "D"]:
            flash("Invalid section selected.", "error")
            return redirect(
                url_for("edit_student", student_id=student_id)
            )

        existing_student = connection.execute("""
            SELECT id
            FROM students
            WHERE student_id = ?
            AND id != ?
        """, (new_student_id, student_id)).fetchone()

        if existing_student:
            connection.close()

            flash(
                "Student ID already exists.",
                "error"
            )

            return redirect(
                url_for(
                    "edit_student",
                    student_id=student_id
                )
            )

        try:

            connection.execute("""
                UPDATE students
                SET
                    student_id = ?,
                    name = ?,
                    email = ?,
                    branch = ?,
                    year = ?,
                    section = ?
                WHERE id = ?
            """, (
                new_student_id,
                name,
                email,
                branch,
                year,
                section,
                student_id
            ))

            connection.commit()
            connection.close()

            flash(
                "Student updated successfully.",
                "success"
            )

            return redirect(url_for("hod_students"))

        except sqlite3.IntegrityError:

            connection.close()

            flash(
                "This email address is already in use.",
                "error"
            )

            return redirect(
                url_for(
                    "edit_student",
                    student_id=student_id
                )
            )

    connection.close()

    return render_template(
        "hod/edit_student.html",
        student=student
    )

@app.route("/hod/students/delete/<int:student_id>", methods=["POST"])
def delete_student(student_id):

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    student = connection.execute("""
        SELECT id, name
        FROM students
        WHERE id = ?
    """, (student_id,)).fetchone()

    if not student:
        connection.close()

        flash(
            "Student not found.",
            "error"
        )

        return redirect(url_for("hod_students"))

    try:

        # Delete attendance records belonging to this student
        connection.execute("""
            DELETE FROM attendance
            WHERE student_id = ?
        """, (student_id,))

        # Delete the student
        connection.execute("""
            DELETE FROM students
            WHERE id = ?
        """, (student_id,))

        connection.commit()
        connection.close()

        flash(
            f"Student '{student['name']}' deleted successfully.",
            "success"
        )

    except sqlite3.Error:

        connection.rollback()
        connection.close()

        flash(
            "Unable to delete the student.",
            "error"
        )

    return redirect(url_for("hod_students"))

@app.route("/hod/faculty")
def hod_faculty():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    search = request.args.get("search", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    if search:

        search_pattern = f"%{search}%"

        faculty = connection.execute("""
            SELECT
                id,
                faculty_id,
                name,
                email
            FROM faculty
            WHERE
                faculty_id LIKE ?
                OR name LIKE ?
                OR email LIKE ?
            ORDER BY id DESC
        """, (
            search_pattern,
            search_pattern,
            search_pattern
        )).fetchall()

    else:

        faculty = connection.execute("""
            SELECT
                id,
                faculty_id,
                name,
                email
            FROM faculty
            ORDER BY id DESC
        """).fetchall()

    connection.close()

    return render_template(
        "hod/faculty.html",
        faculty=faculty,
        search=search
    )

@app.route("/hod/faculty/add", methods=["GET", "POST"])
def add_faculty():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    if request.method == "POST":
        faculty_id = request.form["faculty_id"].strip()
        name = request.form["name"].strip()
        email = request.form["email"].strip()
        username = request.form["username"].strip()
        password = request.form["password"]

        connection = sqlite3.connect(DATABASE_PATH)

        # Check duplicate Faculty ID
        existing_faculty = connection.execute("""
            SELECT id
            FROM faculty
            WHERE faculty_id = ?
        """, (faculty_id,)).fetchone()

        if existing_faculty:
            connection.close()
            flash("Faculty ID already exists.", "error")
            return redirect(url_for("add_faculty"))

        # Check duplicate username
        existing_user = connection.execute("""
            SELECT id
            FROM users
            WHERE username = ?
        """, (username,)).fetchone()

        if existing_user:
            connection.close()
            flash("Username already exists.", "error")
            return redirect(url_for("add_faculty"))

        # Check duplicate email
        existing_email = connection.execute("""
            SELECT id
            FROM faculty
            WHERE email = ?
        """, (email,)).fetchone()

        if existing_email:
            connection.close()
            flash("Email already exists.", "error")
            return redirect(url_for("add_faculty"))

        try:
            # Create login account
            password_hash = generate_password_hash(password)

            cursor = connection.execute("""
                INSERT INTO users (
                    username,
                    password_hash,
                    role
                )
                VALUES (?, ?, 'FACULTY')
            """, (username, password_hash))

            user_id = cursor.lastrowid

            # Create faculty record
            connection.execute("""
                INSERT INTO faculty (
                    faculty_id,
                    name,
                    email,
                    user_id
                )
                VALUES (?, ?, ?, ?)
            """, (
                faculty_id,
                name,
                email,
                user_id
            ))

            connection.commit()
            connection.close()

            flash("Faculty added successfully.", "success")
            return redirect(url_for("hod_faculty"))

        except sqlite3.IntegrityError:
            connection.rollback()
            connection.close()

            flash("Unable to add faculty. Please check the details.", "error")
            return redirect(url_for("add_faculty"))

    return render_template("hod/add_faculty.html")

@app.route("/hod/faculty/edit/<int:faculty_id>", methods=["GET", "POST"])
def edit_faculty(faculty_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    faculty = connection.execute("""
        SELECT
            faculty.id,
            faculty.faculty_id,
            faculty.name,
            faculty.email,
            faculty.user_id,
            users.username
        FROM faculty
        LEFT JOIN users
            ON faculty.user_id = users.id
        WHERE faculty.id = ?
    """, (faculty_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty not found.", "error")
        return redirect(url_for("hod_faculty"))

    if request.method == "POST":

        new_faculty_id = request.form["faculty_id"].strip()
        name = request.form["name"].strip()
        email = request.form["email"].strip()
        username = request.form["username"].strip()

        # Check duplicate Faculty ID
        existing_faculty = connection.execute("""
            SELECT id
            FROM faculty
            WHERE faculty_id = ?
            AND id != ?
        """, (new_faculty_id, faculty_id)).fetchone()

        if existing_faculty:
            connection.close()
            flash("Faculty ID already exists.", "error")
            return redirect(url_for(
                "edit_faculty",
                faculty_id=faculty_id
            ))

        # Check duplicate email
        existing_email = connection.execute("""
            SELECT id
            FROM faculty
            WHERE email = ?
            AND id != ?
        """, (email, faculty_id)).fetchone()

        if existing_email:
            connection.close()
            flash("Email already exists.", "error")
            return redirect(url_for(
                "edit_faculty",
                faculty_id=faculty_id
            ))

        # Check duplicate username
        existing_username = connection.execute("""
            SELECT id
            FROM users
            WHERE username = ?
            AND id != ?
        """, (username, faculty["user_id"])).fetchone()

        if existing_username:
            connection.close()
            flash("Username already exists.", "error")
            return redirect(url_for(
                "edit_faculty",
                faculty_id=faculty_id
            ))

        try:

            # Update faculty information
            connection.execute("""
                UPDATE faculty
                SET faculty_id = ?,
                    name = ?,
                    email = ?
                WHERE id = ?
            """, (
                new_faculty_id,
                name,
                email,
                faculty_id
            ))

            # Update login username
            connection.execute("""
                UPDATE users
                SET username = ?
                WHERE id = ?
            """, (
                username,
                faculty["user_id"]
            ))

            connection.commit()
            connection.close()

            flash("Faculty updated successfully.", "success")

            return redirect(url_for("hod_faculty"))

        except sqlite3.Error:
            connection.rollback()
            connection.close()

            flash(
                "Unable to update faculty.",
                "error"
            )

            return redirect(url_for(
                "edit_faculty",
                faculty_id=faculty_id
            ))

    connection.close()

    return render_template(
        "hod/edit_faculty.html",
        faculty=faculty
    )

@app.route("/hod/faculty/delete/<int:faculty_id>", methods=["POST"])
def delete_faculty(faculty_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    faculty = connection.execute("""
        SELECT id, name, user_id
        FROM faculty
        WHERE id = ?
    """, (faculty_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty not found.", "error")
        return redirect(url_for("hod_faculty"))

    try:
        # Delete faculty record
        connection.execute("""
            DELETE FROM faculty
            WHERE id = ?
        """, (faculty_id,))

        # Delete associated login account
        if faculty["user_id"]:
            connection.execute("""
                DELETE FROM users
                WHERE id = ?
                AND role = 'FACULTY'
            """, (faculty["user_id"],))

        connection.commit()
        connection.close()

        flash(
            f"Faculty '{faculty['name']}' deleted successfully.",
            "success"
        )

    except sqlite3.Error:
        connection.rollback()
        connection.close()

        flash(
            "Unable to delete the faculty.",
            "error"
        )

    return redirect(url_for("hod_faculty"))

@app.route("/hod/subjects")
def hod_subjects():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    search = request.args.get("search", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    if search:
        search_pattern = f"%{search}%"

        subjects = connection.execute("""
            SELECT
                subjects.id,
                subjects.subject_code,
                subjects.subject_name,
                subjects.year,
                subjects.section,
                faculty.name AS faculty_name
            FROM subjects
            LEFT JOIN faculty
                ON subjects.faculty_id = faculty.id
            WHERE subjects.subject_code LIKE ?
               OR subjects.subject_name LIKE ?
               OR faculty.name LIKE ?
               OR subjects.year LIKE ?
               OR subjects.section LIKE ?
            ORDER BY subjects.id DESC
        """, (
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern
        )).fetchall()

    else:

        subjects = connection.execute("""
            SELECT
                subjects.id,
                subjects.subject_code,
                subjects.subject_name,
                subjects.year,
                subjects.section,
                faculty.name AS faculty_name
            FROM subjects
            LEFT JOIN faculty
                ON subjects.faculty_id = faculty.id
            ORDER BY subjects.id DESC
        """).fetchall()

    connection.close()

    return render_template(
        "hod/subjects.html",
        subjects=subjects,
        search=search
    )

@app.route("/hod/subjects/add", methods=["GET", "POST"])
def add_subject():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    # Get faculty members for dropdown
    faculty = connection.execute("""
        SELECT id, faculty_id, name
        FROM faculty
        ORDER BY name ASC
    """).fetchall()

    if request.method == "POST":

        subject_code = request.form["subject_code"].strip()
        subject_name = request.form["subject_name"].strip()
        faculty_id = request.form["faculty_id"]
        year = request.form["year"]
        section = request.form["section"]

        # Check duplicate subject code
        existing_subject = connection.execute("""
            SELECT id
            FROM subjects
            WHERE subject_code = ?
        """, (subject_code,)).fetchone()

        if existing_subject:
            connection.close()

            flash(
                "Subject code already exists.",
                "error"
            )

            return redirect(url_for("add_subject"))

        try:

            connection.execute("""
                INSERT INTO subjects (
                    subject_code,
                    subject_name,
                    faculty_id,
                    year,
                    section
                )
                VALUES (?, ?, ?, ?, ?)
            """, (
                subject_code,
                subject_name,
                faculty_id,
                year,
                section
            ))

            connection.commit()
            connection.close()

            flash(
                "Subject added successfully.",
                "success"
            )

            return redirect(url_for("hod_subjects"))

        except sqlite3.Error:

            connection.rollback()
            connection.close()

            flash(
                "Unable to add subject.",
                "error"
            )

            return redirect(url_for("add_subject"))

    connection.close()

    return render_template(
        "hod/add_subject.html",
        faculty=faculty
    )

@app.route("/hod/subjects/edit/<int:subject_id>", methods=["GET", "POST"])
def edit_subject(subject_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    # Get subject
    subject = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name,
            faculty_id,
            year,
            section
        FROM subjects
        WHERE id = ?
    """, (subject_id,)).fetchone()

    if not subject:
        connection.close()
        flash("Subject not found.", "error")
        return redirect(url_for("hod_subjects"))

    # Get faculty members for dropdown
    faculty = connection.execute("""
        SELECT id, faculty_id, name
        FROM faculty
        ORDER BY name ASC
    """).fetchall()

    if request.method == "POST":

        new_subject_code = request.form["subject_code"].strip()
        subject_name = request.form["subject_name"].strip()
        faculty_id = request.form["faculty_id"]
        year = request.form["year"]
        section = request.form["section"]

        # Check duplicate subject code
        existing_subject = connection.execute("""
            SELECT id
            FROM subjects
            WHERE subject_code = ?
            AND id != ?
        """, (
            new_subject_code,
            subject_id
        )).fetchone()

        if existing_subject:
            connection.close()

            flash(
                "Subject code already exists.",
                "error"
            )

            return redirect(url_for(
                "edit_subject",
                subject_id=subject_id
            ))

        try:

            connection.execute("""
                UPDATE subjects
                SET subject_code = ?,
                    subject_name = ?,
                    faculty_id = ?,
                    year = ?,
                    section = ?
                WHERE id = ?
            """, (
                new_subject_code,
                subject_name,
                faculty_id,
                year,
                section,
                subject_id
            ))

            connection.commit()
            connection.close()

            flash(
                "Subject updated successfully.",
                "success"
            )

            return redirect(url_for("hod_subjects"))

        except sqlite3.Error:

            connection.rollback()
            connection.close()

            flash(
                "Unable to update subject.",
                "error"
            )

            return redirect(url_for(
                "edit_subject",
                subject_id=subject_id
            ))

    connection.close()

    return render_template(
        "hod/edit_subject.html",
        subject=subject,
        faculty=faculty
    )

@app.route("/hod/subjects/delete/<int:subject_id>", methods=["POST"])
def delete_subject(subject_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    subject = connection.execute("""
        SELECT id, subject_name
        FROM subjects
        WHERE id = ?
    """, (subject_id,)).fetchone()

    if not subject:
        connection.close()
        flash("Subject not found.", "error")
        return redirect(url_for("hod_subjects"))

    try:
        # Delete attendance records related to this subject
        connection.execute("""
            DELETE FROM attendance
            WHERE subject_id = ?
        """, (subject_id,))

        # Delete attendance sessions related to this subject
        connection.execute("""
            DELETE FROM attendance_sessions
            WHERE subject_id = ?
        """, (subject_id,))

        # Delete subject
        connection.execute("""
            DELETE FROM subjects
            WHERE id = ?
        """, (subject_id,))

        connection.commit()
        connection.close()

        flash(
            f"Subject '{subject['subject_name']}' deleted successfully.",
            "success"
        )

    except sqlite3.Error:
        connection.rollback()
        connection.close()

        flash(
            "Unable to delete the subject.",
            "error"
        )

    return redirect(url_for("hod_subjects"))

@app.route("/hod/attendance")
def hod_attendance():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    selected_date = request.args.get("date", "").strip()
    selected_subject = request.args.get("subject_id", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    # Get all subjects for the filter
    subjects = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name
        FROM subjects
        ORDER BY subject_code ASC
    """).fetchall()

    # Attendance records query
    query = """
        SELECT
            attendance.id,
            students.student_id,
            students.name AS student_name,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section,
            attendance.date,
            attendance.period,
            attendance.status,
            faculty.name AS faculty_name
        FROM attendance
        JOIN students
            ON attendance.student_id = students.id
        JOIN subjects
            ON attendance.subject_id = subjects.id
        LEFT JOIN faculty
            ON attendance.marked_by = faculty.id
        WHERE 1 = 1
    """

    parameters = []

    if selected_date:
        query += " AND attendance.date = ?"
        parameters.append(selected_date)

    if selected_subject:
        query += " AND attendance.subject_id = ?"
        parameters.append(selected_subject)

    query += """
        ORDER BY attendance.date DESC,
                 students.student_id ASC
    """

    attendance_records = connection.execute(
        query,
        parameters
    ).fetchall()

    # Daily summary
    summary_query = """
        SELECT
            COUNT(*) AS total,
            SUM(
                CASE
                    WHEN status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present,
            SUM(
                CASE
                    WHEN status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent
        FROM attendance
        WHERE 1 = 1
    """

    summary_parameters = []

    if selected_date:
        summary_query += " AND date = ?"
        summary_parameters.append(selected_date)

    if selected_subject:
        summary_query += " AND subject_id = ?"
        summary_parameters.append(selected_subject)

    summary = connection.execute(
        summary_query,
        summary_parameters
    ).fetchone()

    total_attendance = summary["total"] or 0
    present_count = summary["present"] or 0
    absent_count = summary["absent"] or 0

    if total_attendance > 0:
        attendance_percentage = round(
            (present_count / total_attendance) * 100,
            2
        )
    else:
        attendance_percentage = 0

    # Weekly summary - last 7 days
    from datetime import date, timedelta

    today = date.today()
    week_start = today - timedelta(days=6)

    weekly_query = """
        SELECT
            COUNT(*) AS total,
            SUM(
                CASE
                    WHEN status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present,
            SUM(
                CASE
                    WHEN status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent
        FROM attendance
        WHERE date BETWEEN ? AND ?
    """

    weekly_parameters = [
        week_start.isoformat(),
        today.isoformat()
    ]

    if selected_subject:
        weekly_query += " AND subject_id = ?"
        weekly_parameters.append(selected_subject)

    weekly_summary = connection.execute(
        weekly_query,
        weekly_parameters
    ).fetchone()

    weekly_total = weekly_summary["total"] or 0
    weekly_present = weekly_summary["present"] or 0
    weekly_absent = weekly_summary["absent"] or 0

    if weekly_total > 0:
        weekly_percentage = round(
            (weekly_present / weekly_total) * 100,
            2
        )
    else:
        weekly_percentage = 0

        # Monthly summary - current calendar month
    from datetime import date

    today = date.today()

    month_start = today.replace(day=1)

    monthly_query = """
        SELECT
            COUNT(*) AS total,
            SUM(
                CASE
                    WHEN status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present,
            SUM(
                CASE
                    WHEN status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent
        FROM attendance
        WHERE date >= ?
          AND date <= ?
    """

    monthly_parameters = [
        month_start.isoformat(),
        today.isoformat()
    ]

    if selected_subject:
        monthly_query += " AND subject_id = ?"
        monthly_parameters.append(selected_subject)

    monthly_summary = connection.execute(
        monthly_query,
        monthly_parameters
    ).fetchone()

    monthly_total = monthly_summary["total"] or 0
    monthly_present = monthly_summary["present"] or 0
    monthly_absent = monthly_summary["absent"] or 0

    if monthly_total > 0:
        monthly_percentage = round(
            (monthly_present / monthly_total) * 100,
            2
        )
    else:
        monthly_percentage = 0

        # Student-wise attendance percentage
    student_percentage_query = """
        SELECT
            students.id,
            students.student_id,
            students.name,
            COUNT(attendance.id) AS total_classes,
            SUM(
                CASE
                    WHEN attendance.status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present_classes,
            SUM(
                CASE
                    WHEN attendance.status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent_classes
        FROM students
        LEFT JOIN attendance
            ON students.id = attendance.student_id
    """

    student_percentage_parameters = []

    if selected_subject:
        student_percentage_query += """
            AND attendance.subject_id = ?
        """
        student_percentage_parameters.append(selected_subject)

    student_percentage_query += """
        GROUP BY
            students.id,
            students.student_id,
            students.name
        ORDER BY students.student_id ASC
    """

    student_percentage_records = connection.execute(
        student_percentage_query,
        student_percentage_parameters
    ).fetchall()

    connection.close()

    return render_template(
        "hod/attendance.html",
        subjects=subjects,
        attendance_records=attendance_records,
        selected_date=selected_date,
        selected_subject=selected_subject,

        total_attendance=total_attendance,
        present_count=present_count,
        absent_count=absent_count,
        attendance_percentage=attendance_percentage,

        weekly_total=weekly_total,
        weekly_present=weekly_present,
        weekly_absent=weekly_absent,
        weekly_percentage=weekly_percentage,
        week_start=week_start,
        week_end=today,

        monthly_total=monthly_total,
        monthly_present=monthly_present,
        monthly_absent=monthly_absent,
        monthly_percentage=monthly_percentage,
        month_start=month_start,
        month_end=today,

        student_percentage_records=student_percentage_records
    )

@app.route("/hod/attendance/export")
def hod_attendance_export():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    selected_date = request.args.get("date", "").strip()
    selected_subject = request.args.get("subject_id", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    query = """
        SELECT
            students.student_id,
            students.name AS student_name,
            students.branch,
            students.year,
            students.section,
            subjects.subject_code,
            subjects.subject_name,
            attendance.date,
            attendance.period,
            attendance.status,
            faculty.name AS faculty_name
        FROM attendance

        JOIN students
            ON attendance.student_id = students.id

        JOIN subjects
            ON attendance.subject_id = subjects.id

        LEFT JOIN faculty
            ON attendance.marked_by = faculty.id

        WHERE 1 = 1
    """

    parameters = []

    if selected_date:
        query += " AND attendance.date = ?"
        parameters.append(selected_date)

    if selected_subject:
        query += " AND attendance.subject_id = ?"
        parameters.append(selected_subject)

    query += """
        ORDER BY
            attendance.date DESC,
            students.student_id ASC,
            attendance.period ASC
    """

    attendance_records = connection.execute(
        query,
        parameters
    ).fetchall()

    connection.close()

    output = io.StringIO()

    writer = csv.writer(output)

    writer.writerow([
        "Student ID",
        "Student Name",
        "Branch",
        "Year",
        "Section",
        "Subject Code",
        "Subject Name",
        "Date",
        "Period",
        "Status",
        "Marked By"
    ])

    for record in attendance_records:

        writer.writerow([
            record["student_id"],
            record["student_name"],
            record["branch"],
            record["year"],
            record["section"],
            record["subject_code"],
            record["subject_name"],
            record["date"],
            record["period"],
            record["status"],
            record["faculty_name"] or ""
        ])

    csv_data = output.getvalue()

    output.close()

    filename = "attendance_report.csv"

    if selected_date:
        filename = f"attendance_{selected_date}.csv"

    return Response(
        csv_data,
        mimetype="text/csv",
        headers={
            "Content-Disposition":
                f"attachment; filename={filename}"
        }
    )

@app.route("/hod/attendance/export/excel")
def hod_attendance_export_excel():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    selected_date = request.args.get("date", "").strip()
    selected_subject = request.args.get("subject_id", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    selected_subject_name = ""

    if selected_subject:
        subject = connection.execute(
            """
            SELECT subject_code, subject_name
            FROM subjects
            WHERE id = ?
            """,
            (selected_subject,)
        ).fetchone()

        if subject:
            selected_subject_name = (
                f"{subject['subject_code']} - "
                f"{subject['subject_name']}"
            )

    query = """
        SELECT
            students.student_id,
            students.name AS student_name,
            students.branch,
            students.year,
            students.section,
            subjects.subject_code,
            subjects.subject_name,
            attendance.date,
            attendance.period,
            attendance.status,
            faculty.name AS faculty_name
        FROM attendance

        JOIN students
            ON attendance.student_id = students.id

        JOIN subjects
            ON attendance.subject_id = subjects.id

        LEFT JOIN faculty
            ON attendance.marked_by = faculty.id

        WHERE 1 = 1
    """

    parameters = []

    if selected_date:
        query += " AND attendance.date = ?"
        parameters.append(selected_date)

    if selected_subject:
        query += " AND attendance.subject_id = ?"
        parameters.append(selected_subject)

    query += """
        ORDER BY
            attendance.date DESC,
            students.student_id ASC,
            attendance.period ASC
    """

    attendance_records = connection.execute(
        query,
        parameters
    ).fetchall()

    connection.close()

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Attendance Report"

    # ---------------------------------------------------------
    # Title
    # ---------------------------------------------------------

    worksheet.merge_cells("A1:K1")

    worksheet["A1"] = "Student Attendance Report"
    worksheet["A1"].font = Font(
        bold=True,
        size=16
    )
    worksheet["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    worksheet.row_dimensions[1].height = 25

    # ---------------------------------------------------------
    # Filter Information
    # ---------------------------------------------------------

    filter_text = "All Attendance Records"

    if selected_date and selected_subject:
        filter_text = (
            f"Date: {selected_date} | "
            f"Subject: {selected_subject_name}"
        )

    elif selected_date:
        filter_text = f"Date: {selected_date}"

    elif selected_subject:
        filter_text = f"Subject: {selected_subject_name}"

    worksheet.merge_cells("A2:K2")

    worksheet["A2"] = filter_text
    worksheet["A2"].alignment = Alignment(
        horizontal="center"
    )

    # ---------------------------------------------------------
    # Attendance Summary
    # ---------------------------------------------------------

    total_records = len(attendance_records)

    present_count = sum(
        1
        for record in attendance_records
        if str(record["status"]).strip().lower() == "present"
    )

    absent_count = sum(
        1
        for record in attendance_records
        if str(record["status"]).strip().lower() == "absent"
    )

    attendance_percentage = (
        (present_count / total_records) * 100
        if total_records > 0
        else 0
    )

    worksheet.merge_cells("A3:K3")

    worksheet["A3"] = (
        f"Total Records: {total_records} | "
        f"Present: {present_count} | "
        f"Absent: {absent_count} | "
        f"Attendance: {attendance_percentage:.2f}%"
    )

    worksheet["A3"].font = Font(
        bold=True
    )

    worksheet["A3"].alignment = Alignment(
        horizontal="center"
    )

    # ---------------------------------------------------------
    # Table Header
    # ---------------------------------------------------------

    headers = [
        "Student ID",
        "Student Name",
        "Branch",
        "Year",
        "Section",
        "Subject Code",
        "Subject Name",
        "Date",
        "Period",
        "Status",
        "Marked By"
    ]

    header_row = 5

    for column_number, header in enumerate(
        headers,
        start=1
    ):
        cell = worksheet.cell(
            row=header_row,
            column=column_number
        )

        cell.value = header
        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ---------------------------------------------------------
    # Attendance Data
    # ---------------------------------------------------------

    data_start_row = header_row + 1

    for row_number, record in enumerate(
        attendance_records,
        start=data_start_row
    ):

        values = [
            record["student_id"],
            record["student_name"],
            record["branch"],
            record["year"],
            record["section"],
            record["subject_code"],
            record["subject_name"],
            record["date"],
            record["period"],
            record["status"],
            record["faculty_name"] or ""
        ]

        for column_number, value in enumerate(
            values,
            start=1
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column_number
            )

            cell.value = value

            cell.alignment = Alignment(
                vertical="center"
            )

    # ---------------------------------------------------------
    # Empty Result
    # ---------------------------------------------------------

    if not attendance_records:

        worksheet.merge_cells("A6:K6")

        worksheet["A6"] = (
            "No attendance records found "
            "for the selected filters."
        )

        worksheet["A6"].font = Font(
            bold=True
        )

        worksheet["A6"].alignment = Alignment(
            horizontal="center"
        )

    # ---------------------------------------------------------
    # Column Widths
    # ---------------------------------------------------------

    column_widths = {
        "A": 15,
        "B": 22,
        "C": 15,
        "D": 10,
        "E": 10,
        "F": 16,
        "G": 28,
        "H": 15,
        "I": 10,
        "J": 15,
        "K": 20
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    # ---------------------------------------------------------
    # Freeze Header
    # ---------------------------------------------------------

    worksheet.freeze_panes = "A6"

    # ---------------------------------------------------------
    # Auto Filter
    # ---------------------------------------------------------

    if attendance_records:
        last_row = data_start_row + len(attendance_records) - 1

        worksheet.auto_filter.ref = (
            f"A{header_row}:K{last_row}"
        )

    # ---------------------------------------------------------
    # Save Excel File
    # ---------------------------------------------------------

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = "attendance_report.xlsx"

    if selected_date:
        filename = f"attendance_{selected_date}.xlsx"

    return Response(
        output.getvalue(),
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f"attachment; filename={filename}"
        }
    )

@app.route("/hod/attendance/export/pdf")
def hod_attendance_export_pdf():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    selected_date = request.args.get("date", "").strip()
    selected_subject = request.args.get("subject_id", "").strip()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    if selected_subject:
        subject = connection.execute(
            """
            SELECT subject_code, subject_name
            FROM subjects
            WHERE id = ?
            """,
            (selected_subject,)
        ).fetchone()

        if subject:
            selected_subject_name = (
                f"{subject['subject_code']} - "
                f"{subject['subject_name']}"
            )

    query = """
        SELECT
            students.student_id,
            students.name AS student_name,
            students.branch,
            students.year,
            students.section,
            subjects.subject_code,
            subjects.subject_name,
            attendance.date,
            attendance.period,
            attendance.status,
            faculty.name AS faculty_name
        FROM attendance

        JOIN students
            ON attendance.student_id = students.id

        JOIN subjects
            ON attendance.subject_id = subjects.id

        LEFT JOIN faculty
            ON attendance.marked_by = faculty.id

        WHERE 1 = 1
    """

    parameters = []

    if selected_date:
        query += " AND attendance.date = ?"
        parameters.append(selected_date)

    if selected_subject:
        query += " AND attendance.subject_id = ?"
        parameters.append(selected_subject)

    query += """
        ORDER BY
            attendance.date DESC,
            students.student_id ASC,
            attendance.period ASC
    """

    attendance_records = connection.execute(
        query,
        parameters
    ).fetchall()

    connection.close()

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()

    elements = []

    title = Paragraph(
        "<b>Student Attendance Report</b>",
        styles["Title"]
    )

    elements.append(title)
    elements.append(Spacer(1, 10))

    filter_text = "All Attendance Records"

    if selected_date and selected_subject:
        filter_text = (
            f"Date: {selected_date} | "
            f"Subject: {selected_subject_name}"
        )

    elif selected_date:
        filter_text = f"Date: {selected_date}"

    elif selected_subject:
        filter_text = f"Subject: {selected_subject_name}"

    elements.append(
        Paragraph(filter_text, styles["Normal"])
    )

    total_records = len(attendance_records)

    present_count = sum(
        1
        for record in attendance_records
        if str(record["status"]).strip().lower() == "present"
    )

    absent_count = sum(
        1
        for record in attendance_records
        if str(record["status"]).strip().lower() == "absent"
    )

    attendance_percentage = (
        (present_count / total_records) * 100
        if total_records > 0
        else 0
    )

    summary_text = (
        f"<b>Total Records:</b> {total_records} &nbsp;&nbsp; "
        f"<b>Present:</b> {present_count} &nbsp;&nbsp; "
        f"<b>Absent:</b> {absent_count} &nbsp;&nbsp; "
        f"<b>Attendance:</b> {attendance_percentage:.2f}%"
    )

    elements.append(
        Paragraph(summary_text, styles["Normal"])
    )

    elements.append(Spacer(1, 15))

    if not attendance_records:
        elements.append(
            Paragraph(
                "<b>No attendance records found for the selected filters.</b>",
                styles["Normal"]
            )
        )

    document.build(elements)

    output.seek(0)

    return Response(
        output.getvalue(),
        mimetype="application/pdf",
        headers={
            "Content-Disposition":
                "attachment; filename=attendance_report.pdf"
        }
    )

    table_data = [
        [
            "Student ID",
            "Student Name",
            "Branch",
            "Year",
            "Section",
            "Subject Code",
            "Subject Name",
            "Date",
            "Period",
            "Status",
            "Marked By"
        ]
    ]

    for record in attendance_records:

        table_data.append([
            record["student_id"],
            record["student_name"],
            record["branch"],
            record["year"],
            record["section"],
            record["subject_code"],
            record["subject_name"],
            record["date"],
            record["period"],
            record["status"],
            record["faculty_name"] or ""
        ])

    table = Table(
        table_data,
        repeatRows=1
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.grey
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, 0),
                "CENTER"
            ),
            (
                "ALIGN",
                (0, 1),
                (-1, -1),
                "CENTER"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, -1),
                7
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, 0),
                7
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, 0),
                7
            )
        ])
    )

    elements.append(table)

    document.build(elements)

    output.seek(0)

    filename = "attendance_report.pdf"

    if selected_date:
        filename = f"attendance_{selected_date}.pdf"

    return Response(
        output.getvalue(),
        mimetype="application/pdf",
        headers={
            "Content-Disposition":
                f"attachment; filename={filename}"
        }
    )

@app.route("/hod/reports")
def hod_reports():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "HOD":
        return "Access denied", 403

    from datetime import date, timedelta

    today = date.today()

    from_date = request.args.get("from_date", "").strip()
    to_date = request.args.get("to_date", "").strip()
    selected_subject = request.args.get("subject_id", "").strip()

    # Default report period: last 7 days
    if not from_date:
        from_date = (today - timedelta(days=6)).isoformat()

    if not to_date:
        to_date = today.isoformat()

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    # Subjects for filter
    subjects = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name
        FROM subjects
        ORDER BY subject_code ASC
    """).fetchall()

    # Overall report summary
    summary_query = """
        SELECT
            COUNT(*) AS total,
            SUM(
                CASE
                    WHEN status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present,
            SUM(
                CASE
                    WHEN status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent
        FROM attendance
        WHERE date BETWEEN ? AND ?
    """

    summary_parameters = [
        from_date,
        to_date
    ]

    if selected_subject:
        summary_query += """
            AND subject_id = ?
        """
        summary_parameters.append(selected_subject)

    summary = connection.execute(
        summary_query,
        summary_parameters
    ).fetchone()

    total_records = summary["total"] or 0
    present_count = summary["present"] or 0
    absent_count = summary["absent"] or 0

    if total_records > 0:
        attendance_percentage = round(
            (present_count / total_records) * 100,
            2
        )
    else:
        attendance_percentage = 0

    # Student-wise report
    report_query = """
        SELECT
            students.student_id,
            students.name AS student_name,
            COUNT(attendance.id) AS total_classes,
            SUM(
                CASE
                    WHEN attendance.status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present_count,
            SUM(
                CASE
                    WHEN attendance.status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent_count
        FROM students
        LEFT JOIN attendance
            ON students.id = attendance.student_id
            AND attendance.date BETWEEN ? AND ?
    """

    report_parameters = [
        from_date,
        to_date
    ]

    if selected_subject:
        report_query += """
            AND attendance.subject_id = ?
        """
        report_parameters.append(selected_subject)

    report_query += """
        GROUP BY
            students.id,
            students.student_id,
            students.name
        ORDER BY students.student_id ASC
    """

    report_rows = connection.execute(
        report_query,
        report_parameters
    ).fetchall()

    # Calculate percentage for each student
    report_records = []

    for row in report_rows:

        total_classes = row["total_classes"] or 0
        present = row["present_count"] or 0
        absent = row["absent_count"] or 0

        if total_classes > 0:
            percentage = round(
                (present / total_classes) * 100,
                2
            )
        else:
            percentage = 0

        report_records.append({
            "student_id": row["student_id"],
            "student_name": row["student_name"],
            "total_classes": total_classes,
            "present_count": present,
            "absent_count": absent,
            "percentage": percentage
        })

    connection.close()

    return render_template(
        "hod/reports.html",
        subjects=subjects,
        report_records=report_records,
        total_records=total_records,
        present_count=present_count,
        absent_count=absent_count,
        attendance_percentage=attendance_percentage,
        from_date=from_date,
        to_date=to_date,
        selected_subject=selected_subject
    )

@app.route("/faculty/dashboard")
def faculty_dashboard():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "FACULTY":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    # -----------------------------------------
    # Get logged-in faculty
    # -----------------------------------------

    faculty = connection.execute("""
        SELECT
            id,
            faculty_id,
            name
        FROM faculty
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty profile not found.", "error")
        return redirect(url_for("logout"))

    faculty_id = faculty["id"]


    # -----------------------------------------
    # Total assigned subjects
    # -----------------------------------------

    total_subjects = connection.execute("""
        SELECT COUNT(*)
        FROM subjects
        WHERE faculty_id = ?
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Total attendance sessions
    # -----------------------------------------

    total_sessions = connection.execute("""
        SELECT COUNT(*)
        FROM attendance_sessions
        WHERE faculty_id = ?
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Today's attendance sessions
    # -----------------------------------------

    today_sessions = connection.execute("""
        SELECT COUNT(*)
        FROM attendance_sessions
        WHERE faculty_id = ?
          AND date = DATE('now', 'localtime')
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Total students in assigned classes
    # -----------------------------------------

    total_students = connection.execute("""
        SELECT COUNT(DISTINCT students.id)
        FROM students
        JOIN subjects
            ON students.year = subjects.year
            AND students.section = subjects.section
        WHERE subjects.faculty_id = ?
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Today's attendance percentage
    # -----------------------------------------

    today_attendance = connection.execute("""
        SELECT COUNT(*)
        FROM attendance
        WHERE marked_by = ?
          AND date = DATE('now', 'localtime')
    """, (faculty_id,)).fetchone()[0]

    today_present = connection.execute("""
        SELECT COUNT(*)
        FROM attendance
        WHERE marked_by = ?
          AND date = DATE('now', 'localtime')
          AND status = 'PRESENT'
    """, (faculty_id,)).fetchone()[0]

    if today_attendance > 0:

        today_percentage = round(
            (today_present / today_attendance) * 100,
            1
        )

    else:

        today_percentage = 0


    # -----------------------------------------
    # Total attendance records
    # -----------------------------------------

    total_attendance = connection.execute("""
        SELECT COUNT(*)
        FROM attendance
        WHERE marked_by = ?
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Total present
    # -----------------------------------------

    total_present = connection.execute("""
        SELECT COUNT(*)
        FROM attendance
        WHERE marked_by = ?
          AND status = 'PRESENT'
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Total absent
    # -----------------------------------------

    total_absent = connection.execute("""
        SELECT COUNT(*)
        FROM attendance
        WHERE marked_by = ?
          AND status = 'ABSENT'
    """, (faculty_id,)).fetchone()[0]


    # -----------------------------------------
    # Overall attendance percentage
    # -----------------------------------------

    if total_attendance > 0:

        overall_percentage = round(
            (total_present / total_attendance) * 100,
            1
        )

    else:

        overall_percentage = 0


    # -----------------------------------------
    # Subject-wise attendance
    # -----------------------------------------

    subject_summary = connection.execute("""
        SELECT
            subjects.id,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section,

            COUNT(attendance.id) AS total_records,

            SUM(
                CASE
                    WHEN attendance.status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present_count,

            SUM(
                CASE
                    WHEN attendance.status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent_count

        FROM subjects

        LEFT JOIN attendance
            ON subjects.id = attendance.subject_id
            AND attendance.marked_by = ?

        WHERE subjects.faculty_id = ?

        GROUP BY
            subjects.id,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section

        ORDER BY
            subjects.subject_code ASC

    """, (
        faculty_id,
        faculty_id
    )).fetchall()


    # -----------------------------------------
    # Calculate subject percentages
    # -----------------------------------------

    subject_data = []

    for item in subject_summary:

        total = item["total_records"] or 0
        present = item["present_count"] or 0
        absent = item["absent_count"] or 0

        if total > 0:

            percentage = round(
                (present / total) * 100,
                1
            )

        else:

            percentage = 0

        subject_data.append({
            "id": item["id"],
            "subject_code": item["subject_code"],
            "subject_name": item["subject_name"],
            "year": item["year"],
            "section": item["section"],
            "total_records": total,
            "present_count": present,
            "absent_count": absent,
            "percentage": percentage
        })


    # -----------------------------------------
    # Recent attendance sessions
    # -----------------------------------------

    recent_sessions = connection.execute("""
        SELECT
            attendance_sessions.id,
            attendance_sessions.date,
            attendance_sessions.period,

            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section

        FROM attendance_sessions

        JOIN subjects
            ON attendance_sessions.subject_id = subjects.id

        WHERE attendance_sessions.faculty_id = ?

        ORDER BY
            attendance_sessions.date DESC,
            CAST(attendance_sessions.period AS INTEGER) DESC

        LIMIT 5

    """, (faculty_id,)).fetchall()


    # -----------------------------------------
    # Close database connection
    # -----------------------------------------

    connection.close()


    # -----------------------------------------
    # Faculty dashboard
    # -----------------------------------------

    return render_template(
        "faculty/dashboard.html",

        total_subjects=total_subjects,
        total_students=total_students,

        today_sessions=today_sessions,
        today_attendance=today_percentage,

        total_sessions=total_sessions,
        total_attendance=total_attendance,
        total_present=total_present,
        total_absent=total_absent,
        overall_percentage=overall_percentage,

        subject_summary=subject_data,
        recent_sessions=recent_sessions
    )

@app.route("/faculty/subjects")
def faculty_subjects():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "FACULTY":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    # Find logged-in faculty
    faculty = connection.execute("""
        SELECT id
        FROM faculty
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not faculty:
        connection.close()

        flash(
            "Faculty profile not found.",
            "error"
        )

        return redirect(url_for("logout"))

    faculty_id = faculty["id"]

    # Get subjects assigned to this faculty
    subjects = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name,
            year,
            section
        FROM subjects
        WHERE faculty_id = ?
        ORDER BY year ASC, section ASC, subject_code ASC
    """, (faculty_id,)).fetchall()

    connection.close()

    return render_template(
        "faculty/subjects.html",
        subjects=subjects
    )

@app.route("/faculty/attendance", methods=["GET", "POST"])
def faculty_attendance():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "FACULTY":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    faculty = connection.execute("""
        SELECT id, faculty_id, name
        FROM faculty
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty profile not found.", "error")
        return redirect(url_for("logout"))

    faculty_id = faculty["id"]

    subject_id = request.args.get("subject_id", "").strip()

    if not subject_id:
        connection.close()
        flash("Please select a subject.", "warning")
        return redirect(url_for("faculty_subjects"))

    subject = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name,
            year,
            section
        FROM subjects
        WHERE id = ?
          AND faculty_id = ?
    """, (subject_id, faculty_id)).fetchone()

    if not subject:
        connection.close()
        flash("Subject not found or not assigned to you.", "error")
        return redirect(url_for("faculty_subjects"))

    students = connection.execute("""
        SELECT
            id,
            student_id,
            name,
            email,
            branch,
            year,
            section
        FROM students
        WHERE year = ?
          AND section = ?
        ORDER BY student_id ASC
    """, (subject["year"], subject["section"])).fetchall()

    # -------------------------------------------------
    # SAVE ATTENDANCE
    # -------------------------------------------------

    if request.method == "POST":

        attendance_date = request.form.get("attendance_date", "").strip()
        period = request.form.get("period", "").strip()

        if not attendance_date:
            connection.close()
            flash("Please select the attendance date.", "warning")
            return redirect(
                url_for(
                    "faculty_attendance",
                    subject_id=subject_id
                )
            )

        if not period:
            connection.close()
            flash("Please select a period.", "warning")
            return redirect(
                url_for(
                    "faculty_attendance",
                    subject_id=subject_id
                )
            )

        # Check whether attendance session already exists
        existing_session = connection.execute("""
            SELECT id
            FROM attendance_sessions
            WHERE subject_id = ?
              AND faculty_id = ?
              AND date = ?
              AND period = ?
        """, (
            subject_id,
            faculty_id,
            attendance_date,
            period
        )).fetchone()

        if existing_session:
            connection.close()
            flash(
                "Attendance for this subject, date and period already exists.",
                "warning"
            )
            return redirect(
                url_for(
                    "faculty_attendance",
                    subject_id=subject_id
                )
            )

        # Validate attendance for every student
        attendance_records = []

        for student in students:

            status = request.form.get(
                f"attendance_{student['id']}"
            )

            if status not in ("PRESENT", "ABSENT"):
                connection.close()
                flash(
                    f"Please mark attendance for {student['name']}.",
                    "warning"
                )
                return redirect(
                    url_for(
                        "faculty_attendance",
                        subject_id=subject_id
                    )
                )

            attendance_records.append(
                (
                    student["id"],
                    subject_id,
                    attendance_date,
                    period,
                    status,
                    faculty_id
                )
            )

        try:

            # Create attendance session
            cursor = connection.execute("""
                INSERT INTO attendance_sessions
                (
                    subject_id,
                    faculty_id,
                    date,
                    period
                )
                VALUES (?, ?, ?, ?)
            """, (
                subject_id,
                faculty_id,
                attendance_date,
                period
            ))

            # Save each student's attendance
            connection.executemany("""
                INSERT INTO attendance
                (
                    student_id,
                    subject_id,
                    date,
                    period,
                    status,
                    marked_by
                )
                VALUES (?, ?, ?, ?, ?, ?)
            """, attendance_records)

            connection.commit()

            flash(
                "Attendance saved successfully.",
                "success"
            )

        except sqlite3.IntegrityError:

            connection.rollback()

            flash(
                "Attendance could not be saved. "
                "A record may already exist for this date.",
                "error"
            )

        finally:
            connection.close()

        return redirect(
            url_for(
                "faculty_attendance",
                subject_id=subject_id
            )
        )

    connection.close()

    return render_template(
        "faculty/attendance.html",
        subject=subject,
        students=students
    )

@app.route("/faculty/history")
def faculty_history():
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "FACULTY":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    faculty = connection.execute("""
        SELECT id
        FROM faculty
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty profile not found.", "error")
        return redirect(url_for("logout"))

    faculty_id = faculty["id"]

    # Filters
    selected_subject = request.args.get("subject_id", "").strip()
    from_date = request.args.get("from_date", "").strip()
    to_date = request.args.get("to_date", "").strip()

    # Get faculty subjects
    subjects = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name,
            year,
            section
        FROM subjects
        WHERE faculty_id = ?
        ORDER BY subject_code ASC
    """, (faculty_id,)).fetchall()

    # Build session query
    query = """
        SELECT
            attendance_sessions.id,
            attendance_sessions.date,
            attendance_sessions.period,
            subjects.id AS subject_id,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section,

            COUNT(attendance.id) AS total_students,

            SUM(
                CASE
                    WHEN attendance.status = 'PRESENT'
                    THEN 1
                    ELSE 0
                END
            ) AS present_count,

            SUM(
                CASE
                    WHEN attendance.status = 'ABSENT'
                    THEN 1
                    ELSE 0
                END
            ) AS absent_count

        FROM attendance_sessions

        JOIN subjects
            ON attendance_sessions.subject_id = subjects.id

        LEFT JOIN attendance
            ON attendance.subject_id = attendance_sessions.subject_id
            AND attendance.date = attendance_sessions.date
            AND attendance.period = attendance_sessions.period

        WHERE attendance_sessions.faculty_id = ?
    """

    parameters = [faculty_id]

    if selected_subject:
        query += """
            AND attendance_sessions.subject_id = ?
        """
        parameters.append(selected_subject)

    if from_date:
        query += """
            AND attendance_sessions.date >= ?
        """
        parameters.append(from_date)

    if to_date:
        query += """
            AND attendance_sessions.date <= ?
        """
        parameters.append(to_date)

    query += """
        GROUP BY
            attendance_sessions.id,
            attendance_sessions.date,
            attendance_sessions.period,
            subjects.id,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section

        ORDER BY
            attendance_sessions.date DESC,
            CAST(attendance_sessions.period AS INTEGER) DESC
    """

    sessions = connection.execute(
        query,
        parameters
    ).fetchall()

    connection.close()

    # Calculate percentage for each session
    session_data = []

    for item in sessions:

        total = item["total_students"] or 0
        present = item["present_count"] or 0
        absent = item["absent_count"] or 0

        if total > 0:
            percentage = round(
                (present / total) * 100,
                1
            )
        else:
            percentage = 0

        session_data.append({
            "id": item["id"],
            "date": item["date"],
            "period": item["period"],
            "subject_id": item["subject_id"],
            "subject_code": item["subject_code"],
            "subject_name": item["subject_name"],
            "year": item["year"],
            "section": item["section"],
            "total_students": total,
            "present_count": present,
            "absent_count": absent,
            "percentage": percentage
        })

    return render_template(
        "faculty/history.html",
        sessions=session_data,
        subjects=subjects,
        selected_subject=selected_subject,
        from_date=from_date,
        to_date=to_date
    )

@app.route("/faculty/attendance/edit/<int:session_id>", methods=["GET", "POST"])
def faculty_edit_attendance(session_id):
    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "FACULTY":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    faculty = connection.execute("""
        SELECT id
        FROM faculty
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not faculty:
        connection.close()
        flash("Faculty profile not found.", "error")
        return redirect(url_for("logout"))

    faculty_id = faculty["id"]

    attendance_session = connection.execute("""
        SELECT
            attendance_sessions.id,
            attendance_sessions.subject_id,
            attendance_sessions.faculty_id,
            attendance_sessions.date,
            attendance_sessions.period,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section
        FROM attendance_sessions
        JOIN subjects
            ON attendance_sessions.subject_id = subjects.id
        WHERE attendance_sessions.id = ?
          AND attendance_sessions.faculty_id = ?
    """, (session_id, faculty_id)).fetchone()

    if not attendance_session:
        connection.close()
        flash("Attendance session not found.", "error")
        return redirect(url_for("faculty_history"))

    students = connection.execute("""
        SELECT
            students.id,
            students.student_id,
            students.name,
            students.branch,
            students.year,
            students.section,
            attendance.status
        FROM students
        LEFT JOIN attendance
            ON students.id = attendance.student_id
            AND attendance.subject_id = ?
            AND attendance.date = ?
            AND attendance.period = ?
        WHERE students.year = ?
          AND students.section = ?
        ORDER BY students.student_id ASC
    """, (
        attendance_session["subject_id"],
        attendance_session["date"],
        attendance_session["period"],
        attendance_session["year"],
        attendance_session["section"]
    )).fetchall()

    if request.method == "POST":

        attendance_records = []

        for student in students:

            status = request.form.get(
                f"attendance_{student['id']}"
            )

            if status not in ("PRESENT", "ABSENT"):
                connection.close()

                flash(
                    f"Please mark attendance for {student['name']}.",
                    "warning"
                )

                return redirect(
                    url_for(
                        "faculty_edit_attendance",
                        session_id=session_id
                    )
                )

            attendance_records.append(
                (
                    status,
                    student["id"],
                    attendance_session["subject_id"],
                    attendance_session["date"],
                    attendance_session["period"]
                )
            )

        try:

            for status, student_id, subject_id, attendance_date, period in attendance_records:

                existing_record = connection.execute("""
                    SELECT id
                    FROM attendance
                    WHERE student_id = ?
                      AND subject_id = ?
                      AND date = ?
                      AND period = ?
                """, (
                    student_id,
                    subject_id,
                    attendance_date,
                    period
                )).fetchone()

                if existing_record:

                    connection.execute("""
                        UPDATE attendance
                        SET status = ?,
                            marked_by = ?
                        WHERE id = ?
                    """, (
                        status,
                        faculty_id,
                        existing_record["id"]
                    ))

                else:

                    connection.execute("""
                        INSERT INTO attendance
                        (
                            student_id,
                            subject_id,
                            date,
                            period,
                            status,
                            marked_by
                        )
                        VALUES (?, ?, ?, ?, ?, ?)
                    """, (
                        student_id,
                        subject_id,
                        attendance_date,
                        attendance_session["period"],
                        status,
                        faculty_id
                    ))

            connection.commit()

            connection.close()

            flash(
                "Attendance updated successfully.",
                "success"
            )

            return redirect(url_for("faculty_history"))

        except sqlite3.Error:

            connection.rollback()
            connection.close()

            flash(
                "Attendance could not be updated.",
                "error"
            )

            return redirect(
                url_for(
                    "faculty_edit_attendance",
                    session_id=session_id
                )
            )

    connection.close()

    return render_template(
        "faculty/edit_attendance.html",
        attendance_session=attendance_session,
        students=students
    )

@app.route("/student/dashboard")
def student_dashboard():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "STUDENT":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    # Find the logged-in student's profile
    student = connection.execute("""
        SELECT
            id,
            student_id,
            name,
            email,
            branch,
            year,
            section
        FROM students
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not student:
        connection.close()

        flash(
            "Student profile not found.",
            "error"
        )

        return redirect(url_for("logout"))

    student_db_id = student["id"]

    # Overall attendance
    attendance_summary = connection.execute("""
        SELECT
            COUNT(*) AS total_classes,
            SUM(
                CASE
                    WHEN status = 'PRESENT' THEN 1
                    ELSE 0
                END
            ) AS present_classes,
            SUM(
                CASE
                    WHEN status = 'ABSENT' THEN 1
                    ELSE 0
                END
            ) AS absent_classes
        FROM attendance
        WHERE student_id = ?
    """, (student_db_id,)).fetchone()

    total_classes = attendance_summary["total_classes"] or 0
    present_classes = attendance_summary["present_classes"] or 0
    absent_classes = attendance_summary["absent_classes"] or 0

    if total_classes > 0:
        attendance_percentage = round(
            (present_classes / total_classes) * 100,
            2
        )
    else:
        attendance_percentage = 0

    # Subject-wise attendance
    subject_attendance = connection.execute("""
        SELECT
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section,
            COUNT(attendance.id) AS total_classes,
            SUM(
                CASE
                    WHEN attendance.status = 'PRESENT' THEN 1
                    ELSE 0
                END
            ) AS present_classes,
            SUM(
                CASE
                    WHEN attendance.status = 'ABSENT' THEN 1
                    ELSE 0
                END
            ) AS absent_classes
        FROM subjects
        LEFT JOIN attendance
            ON subjects.id = attendance.subject_id
            AND attendance.student_id = ?
        WHERE subjects.year = ?
          AND subjects.section = ?
        GROUP BY
            subjects.id,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section
        ORDER BY subjects.subject_code ASC
    """, (
        student_db_id,
        student["year"],
        student["section"]
    )).fetchall()

    # Recent attendance history
    attendance_history = connection.execute("""
        SELECT
            attendance.date,
            attendance.period,
            attendance.status,
            subjects.subject_code,
            subjects.subject_name
        FROM attendance
        JOIN subjects
            ON attendance.subject_id = subjects.id
        WHERE attendance.student_id = ?
        ORDER BY
            attendance.date DESC,
            attendance.period DESC
        LIMIT 20
    """, (student_db_id,)).fetchall()

    connection.close()

    return render_template(
        "student/dashboard.html",
        student=student,
        total_classes=total_classes,
        present_classes=present_classes,
        absent_classes=absent_classes,
        attendance_percentage=attendance_percentage,
        subject_attendance=subject_attendance,
        attendance_history=attendance_history
    )

@app.route("/student/history")
def student_history():

    if "user_id" not in session:
        return redirect(url_for("login"))

    if session.get("role") != "STUDENT":
        return "Access denied", 403

    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row

    user_id = session["user_id"]

    # Get logged-in student's profile
    student = connection.execute("""
        SELECT
            id,
            student_id,
            name,
            email,
            branch,
            year,
            section
        FROM students
        WHERE user_id = ?
    """, (user_id,)).fetchone()

    if not student:
        connection.close()

        flash(
            "Student profile not found.",
            "error"
        )

        return redirect(url_for("logout"))

    # Filter values
    selected_subject = request.args.get(
        "subject_id",
        ""
    ).strip()

    from_date = request.args.get(
        "from_date",
        ""
    ).strip()

    to_date = request.args.get(
        "to_date",
        ""
    ).strip()

    # Get subjects available for this student's class
    subjects = connection.execute("""
        SELECT
            id,
            subject_code,
            subject_name
        FROM subjects
        WHERE year = ?
          AND section = ?
        ORDER BY subject_code ASC
    """, (
        student["year"],
        student["section"]
    )).fetchall()

    # Attendance history query
    query = """
        SELECT
            attendance.date,
            attendance.period,
            attendance.status,
            subjects.subject_code,
            subjects.subject_name,
            subjects.year,
            subjects.section
        FROM attendance
        JOIN subjects
            ON attendance.subject_id = subjects.id
        WHERE attendance.student_id = ?
    """

    parameters = [
        student["id"]
    ]

    # Subject filter
    if selected_subject:
        query += """
            AND attendance.subject_id = ?
        """

        parameters.append(
            selected_subject
        )

    # From date filter
    if from_date:
        query += """
            AND attendance.date >= ?
        """

        parameters.append(
            from_date
        )

    # To date filter
    if to_date:
        query += """
            AND attendance.date <= ?
        """

        parameters.append(
            to_date
        )

    query += """
        ORDER BY
            attendance.date DESC,
            attendance.period DESC
    """

    attendance_history = connection.execute(
        query,
        parameters
    ).fetchall()

    connection.close()

    return render_template(
        "student/history.html",
        student=student,
        subjects=subjects,
        attendance_history=attendance_history,
        selected_subject=selected_subject,
        from_date=from_date,
        to_date=to_date
    )


@app.route("/logout")
def logout():

    session.clear()

    return redirect(url_for("login"))


if __name__ == "__main__":
    app.run()